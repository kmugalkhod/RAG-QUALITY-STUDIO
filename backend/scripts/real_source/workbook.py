"""Build the spec 0006 workbook from saved results (run on a machine with openpyxl).

    python workbook.py <results dir> <output.xlsx>

<results dir> holds state.json, structure.json, answers_<round>.jsonl,
manual_scores.csv, verdict.json and findings.json. Summary percentages are Excel
formulas over the Retrieval sheet, so they follow any correction to that sheet.
"""

from __future__ import annotations

import csv
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from run import ROUNDS, load_questions, load_sources, track_of  # noqa: E402

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="555555")
# Rounds that build each track; other rounds reuse BASELINE[track].
TRACK_ROUNDS = {
    "web": [r for r, spec in ROUNDS.items() if "web" in spec["tracks"]],
    "files": [r for r, spec in ROUNDS.items() if "files" in spec["tracks"]],
    "new": [r for r, spec in ROUNDS.items() if "new" in spec["tracks"]],
    "d2": [r for r, spec in ROUNDS.items() if "d2" in spec["tracks"]],
}
GOOD_LABELS = ("correct", "correct-abstain")


def style_sheet(sheet, widths: dict[str, int], header_row: int = 1) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.row == header_row,
                             color="FFFFFF" if cell.row == header_row else "000000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == header_row:
                cell.fill = HEADER_FILL
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)


def add_rows(sheet, header: list[str], rows: list[list]) -> None:
    sheet.append(header)
    for row in rows:
        # Extracted text can hold control characters that a worksheet cell rejects.
        sheet.append([
            ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value
            for value in row
        ])


def main(results: Path, output: Path) -> None:
    state = json.loads((results / "state.json").read_text(encoding="utf-8"))
    structure = json.loads((results / "structure.json").read_text(encoding="utf-8"))
    verdict = json.loads((results / "verdict.json").read_text(encoding="utf-8"))
    findings = json.loads((results / "findings.json").read_text(encoding="utf-8"))
    sources = load_sources()
    questions = {q["id"]: q for q in load_questions()}
    scores = {}
    scores_path = results / "manual_scores.csv"
    if scores_path.exists():
        with scores_path.open(encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                scores[(row["round"], row["question_id"])] = row

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    # Sources ---------------------------------------------------------------
    sheet = workbook.create_sheet("Sources")
    rows = []
    for source_id, config in sources["websites"].items():
        note = sources["website_notes"][source_id]
        selection = config["selection"]
        scope = (
            f"Crawl from {selection['start_url']}, depth {config['max_depth']}, "
            f"at most {config['max_pages']} pages"
            if selection["mode"] == "crawl"
            else f"{len(selection['urls'])} listed URLs"
        )
        rows.append([source_id, "Website", note["name"],
                     selection.get("start_url") or "\n".join(selection.get("urls", [])),
                     note["licence"], scope, sources["fetched_on"], ""])
    for source_id, file in {
        **sources["files"],
        **sources.get("files_v4", {}),
        **sources.get("files_v5", {}),
    }.items():
        kind = file["filename"].rsplit(".", 1)[-1].upper()
        scope = f"{file['pages']} pages" if file.get("pages") else ""
        rows.append([source_id, f"File ({kind})", file["name"], file["url"], file["licence"],
                     scope, file.get("fetched_on", sources["fetched_on"]), file["sha256"]])
    add_rows(sheet, ["ID", "Type", "Name", "URL(s)", "Licence", "Scope", "Fetched on",
                     "SHA-256 (files)"], rows)
    style_sheet(sheet, {"A": 6, "B": 11, "C": 34, "D": 60, "E": 34, "F": 30, "G": 12, "H": 30})

    # Runs ------------------------------------------------------------------
    sheet = workbook.create_sheet("Runs")
    header = ["Round", "Track", "What changed", "Extractor", "OCR", "OCR DPI",
              "Max OCR pages", "Table evidence", "If a file can't be read well",
              "Chunk target / max / overlap", "Website pages", "Status", "Error",
              "Items ready", "Items failed", "Items excluded", "Chunks",
              "Duration (s)", "Embedding cost est. (USD)",
              "Embedding tokens (app-reported)", "Embedding cost (app-reported, USD)"]
    rows = []
    for key, run in state.get("runs", {}).items():
        track, round_id = key.split(":")
        spec = ROUNDS[round_id]
        ocr = {"mode": "auto", "dpi": 200, "max_pages": 50, **spec.get("ocr", {})}
        extract = {"config_version": "layout-ocr-v2", "tables": "preserve",
                   **spec.get("extract", {})}
        chunk = {"target_tokens": 600, "maximum_tokens": 800, "overlap_tokens": 80,
                 **spec.get("chunk", {})}
        quality = spec.get("quality", {}).get(track, "default-v1")
        quality_label = {"default-v1": "Stop and let me review",
                         "warn-v1": "Publish the other files and show warnings"}[quality]
        items = run.get("items", [])
        started = run.get("started_at")
        finished = run.get("finished_at")
        duration = None
        if started and finished:
            duration = round(
                (datetime.fromisoformat(finished.replace("Z", "+00:00")) - datetime.fromisoformat(started.replace("Z", "+00:00"))).total_seconds()
            )
        rows.append([
            round_id, track, spec["label"],
            extract["config_version"] if track != "web" else "n/a (HTML)",
            ocr["mode"] if track != "web" else "n/a",
            ocr["dpi"] if track != "web" else "n/a",
            ocr["max_pages"] if track != "web" else "n/a",
            extract["tables"] if track != "web" else "n/a (HTML)",
            quality_label,
            f"{chunk['target_tokens']} / {chunk['maximum_tokens']} / {chunk['overlap_tokens']}",
            ("Fetched live" if round_id == "R1" else "Reused R1 snapshot") if track == "web" else "n/a",
            run.get("status"), run.get("error") or "",
            sum(i.get("status") in {"ready", "succeeded"} for i in items),
            sum(i.get("status") == "failed" for i in items),
            sum(i.get("status") == "excluded" for i in items),
            run.get("chunk_count"), duration, run.get("embedding_estimate_usd"),
            *app_usage(run),
        ])
    rows.sort(key=lambda r: (list(ROUNDS).index(r[0]), r[1]))
    add_rows(sheet, header, rows)
    last = len(rows) + 1
    sheet.append([])
    sheet.append(["Total embedding estimate (USD)", "", "", "", "", "", "", "", "", "", "",
                  "", "", "", "", "", "", "", f"=SUM(S2:S{last})"])
    sheet.append(["Estimate basis: indexed UTF-8 bytes x $0.02 per million (OpenRouter list "
                  "price for openai/text-embedding-3-small, 2026-10-09). Bytes overstate "
                  "tokens, so this is an upper bound. From A1 (spec 0007) the app records the "
                  "provider-reported usage; earlier runs show 'not recorded'."])
    style_sheet(sheet, {get_column_letter(i): w for i, w in enumerate(
        [7, 7, 30, 14, 7, 8, 9, 13, 24, 14, 16, 10, 40, 8, 8, 8, 8, 9, 12, 12, 12],
        start=1)})
    for row in sheet.iter_rows(min_row=last + 2):
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.row == last + 2)
    sheet.cell(row=last + 3, column=1).font = NOTE_FONT
    runs_total_cell = f"Runs!S{last + 2}"

    # Structure ---------------------------------------------------------------
    sheet = workbook.create_sheet("Structure")
    header = ["Round", "Track", "Source", "Run status", "Items", "Items ready",
              "Items failed", "Items excluded", "Pages or URLs indexed", "Chunks",
              "Median chunk tokens", "Max chunk tokens", "Chunks with section path (%)",
              "Table chunks (pipe rows)", "Table chunks with header",
              "Short single-line chunks"]
    rows = [[s["round"], s["track"], s["source"], s["run_status"], s["items"],
             s["items_ready"], s["items_failed"], s["items_excluded"],
             s["pages_or_urls_indexed"], s["chunks"], s["chunk_tokens_median"],
             s["chunk_tokens_max"], s["chunks_with_section_percent"], s["table_chunks"],
             s["table_chunks_with_header"], s["heading_only_chunks"]] for s in structure]
    rows.sort(key=lambda r: (list(ROUNDS).index(r[0]), r[1], r[2]))
    add_rows(sheet, header, rows)
    style_sheet(sheet, {get_column_letter(i): 11 for i in range(1, 17)})

    # Retrieval -----------------------------------------------------------------
    sheet = workbook.create_sheet("Retrieval")
    header = ["Round", "Source", "Question ID", "Question", "Kind", "Expected phrase(s)",
              "Index used (round)", "First hit rank", "Hit@1", "Hit@5",
              "Top result source", "Top result section", "Top result excerpt"]
    rows = []
    for round_id in ROUNDS:
        results_round = state.get("retrieval", {}).get(round_id, {})
        for question_id, question in questions.items():
            track = track_of(question["source"], sources)
            if track not in ROUNDS[round_id]["tracks"] and round_id != "R1":
                continue
            result = results_round.get(question_id)
            if result is None:
                continue
            top = (result.get("top") or [{}])[0]
            answerable = bool(question["phrases"])
            hit1 = result.get("hit_at_1", False)
            hit5 = result.get("hit_at_5", False)
            rows.append([
                round_id, question["source"], question_id, question["question"],
                question["kind"], " + ".join(question["phrases"]) or "(not in source)",
                result.get("index_round") if result.get("status") == "ok" else "no index",
                result.get("first_hit_rank"),
                ("yes" if hit1 else "no") if answerable else "n/a",
                ("yes" if hit5 else "no") if answerable else "n/a",
                top.get("source"), top.get("section"),
                (top.get("text") or "")[:250],
            ])
    add_rows(sheet, header, rows)
    retrieval_last = len(rows) + 1
    style_sheet(sheet, {"A": 7, "B": 7, "C": 8, "D": 44, "E": 11, "F": 26, "G": 10,
                        "H": 8, "I": 7, "J": 7, "K": 34, "L": 30, "M": 60})

    # Answers -------------------------------------------------------------------
    sheet = workbook.create_sheet("Answers")
    header = ["Round", "Index used", "Question ID", "Question", "Reference answer",
              "Model answer", "Label (blind)", "Reason", "Supported by cited evidence",
              "Supported reason", "Cost (USD)"]
    rows = []
    answer_total = 0.0
    for round_id in ROUNDS:
        path = results / f"answers_{round_id}.jsonl"
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            answer = json.loads(line)
            score = scores.get((round_id, answer["question_id"]), {})
            answer_total += answer.get("cost_usd") or 0
            rows.append([
                round_id, answer.get("index_round"), answer["question_id"],
                questions[answer["question_id"]]["question"],
                questions[answer["question_id"]]["answer"],
                answer.get("answer") or answer.get("status"),
                score.get("label", ""), score.get("reason", ""),
                score.get("supported", ""), score.get("supported_reason", ""),
                answer.get("cost_usd"),
            ])
    add_rows(sheet, header, rows)
    answers_last = len(rows) + 1
    sheet.append([])
    sheet.append(["Total answer cost (USD)", "", "", "", "", "", "", "", "", "",
                  f"=SUM(K2:K{answers_last})"])
    sheet.append(["Labels: correct, correct-abstain, partial, wrong, wrong-abstain. Scored "
                  "blind by Claude with the spec 0004 rubric; rounds were hidden as X/Y. A1 (after "
                  "the fixes) and A2 (extractor v4) were scored with the same rubric against the "
                  "references, not blind."])
    style_sheet(sheet, {"A": 7, "B": 8, "C": 8, "D": 40, "E": 30, "F": 50, "G": 13,
                        "H": 40, "I": 10, "J": 36, "K": 9})
    sheet.cell(row=answers_last + 3, column=1).font = NOTE_FONT
    answers_total_cell = f"Answers!K{answers_last + 2}"

    # Settings verdict ------------------------------------------------------------
    sheet = workbook.create_sheet("Settings verdict")
    add_rows(sheet, ["Setting", "Where it shows", "Default", "Tested in", "Measured effect",
                     "Verdict", "Recommendation"],
             [[v["setting"], v["where"], v["default"], v["tested"], v["effect"],
               v["verdict"], v["recommendation"]] for v in verdict])
    style_sheet(sheet, {"A": 26, "B": 18, "C": 18, "D": 12, "E": 60, "F": 16, "G": 50})

    # Findings ----------------------------------------------------------------------
    sheet = workbook.create_sheet("Findings")
    add_rows(sheet, ["ID", "Severity", "Finding", "Evidence", "Impact on users", "Next step",
                     "After fixes (spec 0007)", "After-fixes evidence"],
             [[f["id"], f["severity"], f["finding"], f["evidence"], f["impact"], f["next"],
               f.get("after_fixes", ""), f.get("after_fixes_evidence", "")]
              for f in findings])
    style_sheet(sheet, {"A": 6, "B": 9, "C": 46, "D": 52, "E": 42, "F": 40, "G": 14, "H": 60})

    # After fixes (spec 0007) ---------------------------------------------------------
    after_fixes_sheet(workbook.create_sheet("After fixes"), state, questions, scores, findings)

    # Extract v4 (spec 0008) ------------------------------------------------------------
    extract_v4_sheet(workbook.create_sheet("Extract v4"), state, questions, scores, sources)

    # Extract v5 (spec 0009) ------------------------------------------------------------
    extract_v5_sheet(workbook.create_sheet("Extract v5"), state, questions, scores, sources)

    # Search v3 (spec 0010) -------------------------------------------------------------
    search_v3_sheet(workbook.create_sheet("Search v3"), state, questions, scores, sources)

    # Summary (formulas over Retrieval and Structure) ----------------------------------
    summary.append(["Real-source ingestion test (spec 0006), re-run after the spec 0007 fixes (A1) "
                    "and with extractor v4 (A2, spec 0008)"])
    summary.append(["Satisfaction rule (fixed before results): Good = Hit@5 at least 90% and the "
                    "run published; Acceptable = Hit@5 at least 75%; Poor = otherwise, or the run "
                    "did not publish."])
    summary.append([])
    header_row = 4
    summary.append(["Round", "What changed", "Source", "Answerable questions", "Hit@1",
                    "Hit@5", "Hit@1 %", "Hit@5 %", "Run published", "Satisfaction"])
    rng = lambda col: f"Retrieval!${col}$2:${col}${retrieval_last}"  # noqa: E731
    row_number = header_row
    for track in ("web", "files", "new", "d2"):
        track_sources = {
            "web": sources["websites"],
            "files": sources["files"],
            "new": sources.get("files_v4", {}),
            "d2": sources.get("files_v5", {}),
        }[track]
        for round_id in TRACK_ROUNDS[track]:
            for source_id in track_sources:
                row_number += 1
                r, s = f"A{row_number}", f"C{row_number}"
                summary.append([
                    round_id, ROUNDS[round_id]["label"], source_id,
                    f'=COUNTIFS({rng("A")},{r},{rng("B")},{s},{rng("I")},"<>n/a")',
                    f'=COUNTIFS({rng("A")},{r},{rng("B")},{s},{rng("I")},"yes")',
                    f'=COUNTIFS({rng("A")},{r},{rng("B")},{s},{rng("J")},"yes")',
                    f"=IF(D{row_number}=0,0,E{row_number}/D{row_number})",
                    f"=IF(D{row_number}=0,0,F{row_number}/D{row_number})",
                    f'=IF(COUNTIFS(Structure!$A:$A,{r},Structure!$C:$C,{s},'
                    f'Structure!$D:$D,"succeeded")>0,"yes","no")',
                    f'=IF(I{row_number}="no","Poor (not published)",IF(H{row_number}>=0.9,"Good",'
                    f'IF(H{row_number}>=0.75,"Acceptable","Poor")))',
                ])
                for column in ("G", "H"):
                    summary[f"{column}{row_number}"].number_format = "0%"
    summary.append([])
    summary.append(["Total embedding estimate (USD)", "", "", f"={runs_total_cell}"])
    summary.append(["Total answer cost (USD)", "", "", f"={answers_total_cell}"])
    summary.append(["Total spend (USD, embedding upper bound + recorded answers)", "", "",
                    f"=D{row_number + 2}+D{row_number + 3}"])
    summary.append(["Spending cap (USD)", "", "", 1])
    summary.append([])
    summary.append(["Recommendation"])
    for line in verdict_summary(verdict, findings):
        summary.append([line])
    style_sheet(summary, {"A": 12, "B": 40, "C": 8, "D": 12, "E": 7, "F": 7, "G": 8,
                          "H": 8, "I": 10, "J": 20}, header_row=header_row)
    summary["A1"].font = Font(name=FONT, size=14, bold=True)
    summary["A2"].font = NOTE_FONT
    summary.freeze_panes = f"A{header_row + 1}"
    workbook.calculation.fullCalcOnLoad = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Wrote {output}")


def app_usage(run: dict) -> list:
    usage = run.get("embedding_usage")
    if usage is None:
        return ["not recorded", "not recorded"]
    return [
        "unknown" if usage.get("tokens") is None else usage["tokens"],
        "unknown" if usage.get("cost_usd") is None else usage["cost_usd"],
    ]


def after_fixes_sheet(sheet, state, questions, scores, findings) -> None:
    """Spec 0007: the recommended settings before (R1 websites, R1w files) and after (A1)."""

    retrieval = state.get("retrieval", {})
    before_round = {"web": "R1", "files": "R1w"}
    sheet.append(["Before vs after the spec 0007 fixes (same website snapshot, files and "
                  "questions)"])
    sheet.append(["Before: websites R1, files R1w (R1 files was blocked). After: A1 = extractor "
                  "v3, files 'Stop and let me review', websites 'Publish the others'."])
    sheet.append([])
    header_row = 4
    sheet.append(["Track", "Measure", "Before", "After", "Note"])
    for track in ("web", "files"):
        ids = [q for q, spec in questions.items()
               if spec["phrases"] and not spec.get("round")
               and (spec["source"].startswith("W") == (track == "web"))]
        before = retrieval.get(before_round[track], {})
        after = retrieval.get("A1", {})
        for measure, key in (("Hit@1", "hit_at_1"), ("Hit@5", "hit_at_5")):
            sheet.append([track, measure,
                          f"{sum(before.get(q, {}).get(key, False) for q in ids)}/{len(ids)}",
                          f"{sum(after.get(q, {}).get(key, False) for q in ids)}/{len(ids)}", ""])
        run_before = state["runs"].get(f"{track}:R1", {})
        run_after = state["runs"].get(f"{track}:A1", {})
        sheet.append([track, "Recommended run published", run_before.get("status"),
                      run_after.get("status"),
                      "R1 files was blocked by the Census report" if track == "files" else ""])
        usage = app_usage(run_after)
        sheet.append([track, "Embedding (app-reported)", "not recorded",
                      f"{usage[0]} tokens, ${usage[1]}", "Reused vectors add nothing"])
    labels = {}
    for (round_id, _), row in scores.items():
        if round_id in ("R1", "R1w", "A1"):
            labels.setdefault(round_id, []).append(row["label"])
    good = ("correct", "correct-abstain")
    before_labels = [row["label"] for (r, q), row in scores.items()
                     if (r == "R1" and q.startswith("w")) or (r == "R1w" and q.startswith("f"))]
    after_labels = labels.get("A1", [])
    sheet.append(["both", "Answers correct or correctly declined",
                  f"{sum(label in good for label in before_labels)}/{len(before_labels)}",
                  f"{sum(label in good for label in after_labels)}/{len(after_labels)}",
                  "A1 scored against the references, not blind (one new round)"])
    sheet.append([])
    sheet.append(["Questions whose retrieval changed"])
    sheet.append(["Question", "Source", "Question text", "First hit rank before",
                  "First hit rank after"])
    for qid, spec in questions.items():
        if not spec["phrases"] or spec.get("round"):
            continue
        track = "web" if spec["source"].startswith("W") else "files"
        old = retrieval.get(before_round[track], {}).get(qid, {}).get("first_hit_rank")
        new = retrieval.get("A1", {}).get(qid, {}).get("first_hit_rank")
        if old != new:
            sheet.append([qid, spec["source"], spec["question"], old or "not in top 5",
                          new or "not in top 5"])
    sheet.append([])
    sheet.append(["Findings status"])
    sheet.append(["ID", "Finding", "After fixes", "Evidence"])
    for finding in findings:
        sheet.append([finding["id"], finding["finding"], finding.get("after_fixes", ""),
                      finding.get("after_fixes_evidence", "")])
    superseded = state.get("superseded", {}).get("web:A1")
    if superseded:
        sheet.append([])
        sheet.append(["Note", superseded["reason"]])
    style_sheet(sheet, {"A": 10, "B": 40, "C": 18, "D": 22, "E": 60}, header_row=header_row)
    sheet["A1"].font = Font(name=FONT, size=14, bold=True)
    sheet["A2"].font = NOTE_FONT


def extract_v4_sheet(sheet, state, questions, scores, sources) -> None:
    """Spec 0008: A1 (extractor v3) against A2 (v4, simplified Extract defaults)."""

    retrieval = state.get("retrieval", {})
    before, after = retrieval.get("A1", {}), retrieval.get("A2", {})
    sheet.append(["Extractor v4 against v3 on real files (spec 0008, round A2)"])
    sheet.append(["Before: A1 = extractor v3, Maximum OCR pages 50. After: A2 = extractor v4, "
                  "Maximum OCR pages 100, other Extract settings at the recommended values. The "
                  "spec 0006 files use 'Stop and let me review'; the new files (F4-F7) use "
                  "'Publish the other files and show warnings'. Websites reuse the A1 index: v4 "
                  "does not change the website reader."])
    sheet.append([])
    header_row = 4
    sheet.append(["Measure", "A1 (v3)", "A2 (v4)", "Note"])
    old_ids = [q for q, spec in questions.items()
               if spec["phrases"] and track_of(spec["source"], sources) == "files"
               and not spec.get("round")]
    for measure, key in (("Hit@1, spec 0006 file questions", "hit_at_1"),
                         ("Hit@5, spec 0006 file questions", "hit_at_5")):
        sheet.append([measure,
                      f"{sum(before.get(q, {}).get(key, False) for q in old_ids)}/{len(old_ids)}",
                      f"{sum(after.get(q, {}).get(key, False) for q in old_ids)}/{len(old_ids)}",
                      ""])
    lost = [q for q in old_ids
            if before.get(q, {}).get("hit_at_5") and not after.get(q, {}).get("hit_at_5")]
    sheet.append(["File questions found in A1 and missed in A2", "", len(lost),
                  ", ".join(lost) or "none"])
    for track in ("files", "new"):
        run_before = state["runs"].get(f"{track}:A1", {})
        run_after = state["runs"].get(f"{track}:A2", {})
        usage = app_usage(run_after)
        sheet.append([f"{track} run", run_before.get("status", "not run"),
                      run_after.get("status", "not run"),
                      f"{run_after.get('chunk_count')} chunks; embedding {usage[0]} tokens, "
                      f"${usage[1]} (app-reported)"])
    old_answers = {q for (r, q) in scores if r == "A1" and q.startswith("f")}
    a1_good = {q for (r, q), row in scores.items() if r == "A1" and row["label"] in GOOD_LABELS}
    a2_good = {q for (r, q), row in scores.items() if r == "A2" and row["label"] in GOOD_LABELS}
    a2_old = {q for (r, q) in scores if r == "A2" and q in old_answers}
    sheet.append(["Answers correct or correctly declined, spec 0006 file questions",
                  f"{len(a1_good & old_answers)}/{len(old_answers)}",
                  f"{len(a2_good & a2_old)}/{len(a2_old)}", ""])
    worse = sorted((a1_good & old_answers) - a2_good)
    sheet.append(["File answers correct in A1 and not in A2", "", len(worse),
                  ", ".join(worse) or "none"])
    new_ids = [q for q, spec in questions.items() if spec.get("round") == "A2"]
    new_answered = [q for q in new_ids if ("A2", q) in scores]
    sheet.append(["Answers correct or correctly declined, new questions", "",
                  f"{len(a2_good & set(new_answered))}/{len(new_answered)}", ""])
    sheet.append([])
    sheet.append(["New questions (written from the source files)"])
    sheet.append(["Question", "Source", "Kind", "Question text", "Reference answer",
                  "First hit rank (A2)", "Answer label (A2)", "Reason"])
    for qid in new_ids:
        spec = questions[qid]
        result = after.get(qid, {})
        rank = result.get("first_hit_rank") if spec["phrases"] else "n/a"
        score = scores.get(("A2", qid), {})
        sheet.append([qid, spec["source"], spec["kind"], spec["question"], spec["answer"],
                      rank or "not in top 5", score.get("label", ""), score.get("reason", "")])
    sheet.append([])
    sheet.append(["Questions whose retrieval changed from A1 to A2"])
    sheet.append(["Question", "Source", "Question text", "First hit rank A1",
                  "First hit rank A2"])
    for qid in old_ids:
        old = before.get(qid, {}).get("first_hit_rank")
        new = after.get(qid, {}).get("first_hit_rank")
        if old != new:
            sheet.append([qid, questions[qid]["source"], questions[qid]["question"],
                          old or "not in top 5", new or "not in top 5"])
    style_sheet(sheet, {"A": 34, "B": 10, "C": 12, "D": 50, "E": 30, "F": 12, "G": 14,
                        "H": 50}, header_row=header_row)
    sheet["A1"].font = Font(name=FONT, size=14, bold=True)
    sheet["A2"].font = NOTE_FONT


def extract_v5_sheet(sheet, state, questions, scores, sources) -> None:
    """Spec 0009: A2 (extractor v4) against A3 (v5), plus the spec 0009 files."""

    retrieval = state.get("retrieval", {})
    before, after = retrieval.get("A2", {}), retrieval.get("A3", {})
    sheet.append(["Extractor v5 against v4 on real files (spec 0009, round A3)"])
    sheet.append(["Before: A2 = extractor v4; the spec 0008 files used 'Publish the other files'. "
                  "After: A3 = extractor v5 with the same settings, the spec 0008 files under "
                  "'Stop and let me review'. The spec 0009 files (G1-G8) use 'Publish the other "
                  "files and show warnings' so each file's outcome shows. Websites reuse the A1 "
                  "index."])
    sheet.append([])
    header_row = 4
    sheet.append(["Measure", "A2 (v4)", "A3 (v5)", "Note"])
    earlier = [q for q, spec in questions.items()
               if spec["phrases"] and track_of(spec["source"], sources) in ("files", "new")
               and spec.get("round") in (None, "A2")]
    for measure, key in (("Hit@1, spec 0006 and 0008 file questions", "hit_at_1"),
                         ("Hit@5, spec 0006 and 0008 file questions", "hit_at_5")):
        sheet.append([measure,
                      f"{sum(before.get(q, {}).get(key, False) for q in earlier)}/{len(earlier)}",
                      f"{sum(after.get(q, {}).get(key, False) for q in earlier)}/{len(earlier)}",
                      ""])
    lost = [q for q in earlier
            if before.get(q, {}).get("hit_at_5") and not after.get(q, {}).get("hit_at_5")]
    sheet.append(["Found in A2 and missed in A3", "", len(lost), ", ".join(lost) or "none"])
    for track in ("files", "new", "d2"):
        run_before = state["runs"].get(f"{track}:A2", {})
        run_after = state["runs"].get(f"{track}:A3", {})
        usage = app_usage(run_after)
        sheet.append([f"{track} run", run_before.get("status", "not run"),
                      run_after.get("status", "not run"),
                      f"{run_after.get('chunk_count')} chunks; embedding {usage[0]} tokens, "
                      f"${usage[1]} (app-reported)"])
    good = lambda round_id: {  # noqa: E731
        q for (r, q), row in scores.items() if r == round_id and row["label"] in GOOD_LABELS
    }
    asked = lambda round_id: {q for (r, q) in scores if r == round_id}  # noqa: E731
    a2_files = {q for q in asked("A2") if q[0] == "f"}
    a3_files = {q for q in asked("A3") if q in a2_files}
    sheet.append(["Answers correct or correctly declined, spec 0006 and 0008 file questions",
                  f"{len(good('A2') & a2_files)}/{len(a2_files)}",
                  f"{len(good('A3') & a3_files)}/{len(a3_files)}", ""])
    worse = sorted((good("A2") & a2_files) - good("A3"))
    better = sorted((good("A3") & a3_files) - good("A2"))
    sheet.append(["File answers correct in A2 and not in A3", "", len(worse),
                  ", ".join(worse) or "none"])
    sheet.append(["File answers correct in A3 and not in A2", "", len(better),
                  ", ".join(better) or "none"])
    new_ids = [q for q, spec in questions.items() if spec.get("round") == "A3"]
    new_answered = [q for q in new_ids if ("A3", q) in scores]
    sheet.append(["Answers correct or correctly declined, spec 0009 questions", "",
                  f"{len(good('A3') & set(new_answered))}/{len(new_answered)}", ""])
    sheet.append([])
    sheet.append(["Spec 0009 files (track d2)"])
    sheet.append(["Source", "File", "Item status", "Outcome or reason", "Chunks"])
    items = {i.get("filename"): i for i in state["runs"].get("d2:A3", {}).get("items", [])}
    for source_id, file in sources.get("files_v5", {}).items():
        item = items.get(file["filename"], {})
        sheet.append([source_id, file["name"], item.get("status", "not run"),
                      item.get("reason") or item.get("outcome") or item.get("error") or "",
                      item.get("chunk_count")])
    sheet.append([])
    sheet.append(["Spec 0009 questions (written from the source files before the run)"])
    sheet.append(["Question", "Source", "Kind", "Question text", "Reference answer",
                  "First hit rank (A3)", "Answer label (A3)", "Reason"])
    for qid in new_ids:
        spec = questions[qid]
        rank = after.get(qid, {}).get("first_hit_rank") if spec["phrases"] else "n/a"
        score = scores.get(("A3", qid), {})
        sheet.append([qid, spec["source"], spec["kind"], spec["question"], spec["answer"],
                      rank or "not in top 5", score.get("label", ""), score.get("reason", "")])
    sheet.append([])
    sheet.append(["Questions whose retrieval changed from A2 to A3"])
    sheet.append(["Question", "Source", "Question text", "First hit rank A2",
                  "First hit rank A3"])
    for qid in earlier:
        old = before.get(qid, {}).get("first_hit_rank")
        new = after.get(qid, {}).get("first_hit_rank")
        if old != new:
            sheet.append([qid, questions[qid]["source"], questions[qid]["question"],
                          old or "not in top 5", new or "not in top 5"])
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
    style_sheet(sheet, {"A": 34, "B": 30, "C": 12, "D": 50, "E": 30, "F": 12, "G": 14,
                        "H": 50}, header_row=header_row)
    sheet["A1"].font = Font(name=FONT, size=14, bold=True)
    sheet["A2"].font = NOTE_FONT


SEARCH_ROUNDS = ("A3", "A4s", "A4p")
# The questions spec 0010 set out to fix, and the other A3 misses.
SEARCH_TARGETS = ("f2-q9", "f2-q10", "f4-q1", "f5-q2", "f5-q3", "f6-q1", "g3-q2")


def search_v3_sheet(sheet, state, questions, scores, sources) -> None:
    """Spec 0010: A3 (section-token-v1) against the v2 chunkers, same files and settings."""

    retrieval = state.get("retrieval", {})
    sheet.append(["Chunking v2 against v1 on real files (spec 0010, round A4)"])
    sheet.append(["A3 = extractor v5 with section-token-v1 (600/800/80). A4s = the same with "
                  "section-token-v2, A4p = parent-child-v2 (240/320/40 children, 900/1200 "
                  "parents), each with the defaults a new pipeline gets. v2 appends a "
                  "section's first chunk under 40 tokens to the chunk before it on the same "
                  "page. Same files, questions, embedding model, vector search top 5 and "
                  "answer model in every round."])
    sheet.append([])
    header_row = 4
    sheet.append(["Measure", *SEARCH_ROUNDS, "Note"])
    answerable = [q for q, spec in questions.items()
                  if spec["phrases"] and track_of(spec["source"], sources) != "web"]
    for measure, key in (("Hit@1, file questions", "hit_at_1"),
                         ("Hit@5, file questions", "hit_at_5")):
        sheet.append([measure, *(
            f"{sum(retrieval.get(r, {}).get(q, {}).get(key, False) for q in answerable)}"
            f"/{len(answerable)}" for r in SEARCH_ROUNDS), ""])
    base = retrieval.get("A3", {})
    for round_id in SEARCH_ROUNDS[1:]:
        after = retrieval.get(round_id, {})
        lost = [q for q in answerable
                if base.get(q, {}).get("hit_at_5") and not after.get(q, {}).get("hit_at_5")]
        gained = [q for q in answerable
                  if after.get(q, {}).get("hit_at_5") and not base.get(q, {}).get("hit_at_5")]
        sheet.append([f"{round_id}: in top 5 in A3 and not now", "", "", "",
                      f"{len(lost)}: {', '.join(lost) or 'none'}"])
        sheet.append([f"{round_id}: in top 5 now and not in A3", "", "", "",
                      f"{len(gained)}: {', '.join(gained) or 'none'}"])
    for track in ("files", "new", "d2"):
        cells = []
        for round_id in SEARCH_ROUNDS:
            run = state["runs"].get(f"{track}:{round_id}", {})
            usage = app_usage(run)
            cells.append(f"{run.get('status', 'not run')}; {run.get('chunk_count')} chunks; "
                         f"{usage[0]} tokens, ${usage[1]}")
        sheet.append([f"{track} run (status; chunks; app-reported embedding)", *cells, ""])
    good = lambda round_id: {  # noqa: E731
        q for (r, q), row in scores.items() if r == round_id and row["label"] in GOOD_LABELS
    }
    asked = lambda round_id: {q for (r, q) in scores if r == round_id}  # noqa: E731
    sheet.append(["Answers correct or correctly declined",
                  *(f"{len(good(r))}/{len(asked(r))}" if asked(r) else "not scored"
                    for r in SEARCH_ROUNDS), ""])
    for round_id in SEARCH_ROUNDS[1:]:
        worse = sorted((good("A3") & asked(round_id)) - good(round_id))
        better = sorted((good(round_id) & asked("A3")) - good("A3"))
        sheet.append([f"{round_id}: answers right in A3 and not now", "", "", "",
                      f"{len(worse)}: {', '.join(worse) or 'none'}"])
        sheet.append([f"{round_id}: answers right now and not in A3", "", "", "",
                      f"{len(better)}: {', '.join(better) or 'none'}"])
    sheet.append([])
    sheet.append(["A3 misses (first hit rank; answer label)"])
    sheet.append(["Question", *SEARCH_ROUNDS, "Question text"])
    for qid in SEARCH_TARGETS:
        cells = []
        for round_id in SEARCH_ROUNDS:
            rank = retrieval.get(round_id, {}).get(qid, {}).get("first_hit_rank")
            label = scores.get((round_id, qid), {}).get("label", "not scored")
            cells.append(f"{rank or 'not in top 5'}; {label}")
        sheet.append([qid, *cells, questions[qid]["question"]])
    sheet.append([])
    sheet.append(["Questions whose first hit rank changed from A3"])
    sheet.append(["Question", *SEARCH_ROUNDS, "Question text"])
    for qid in answerable:
        ranks = [retrieval.get(r, {}).get(qid, {}).get("first_hit_rank") for r in SEARCH_ROUNDS]
        if len(set(ranks)) > 1:
            sheet.append([qid, *(rank or "not in top 5" for rank in ranks),
                          questions[qid]["question"]])
    style_sheet(sheet, {"A": 40, "B": 26, "C": 26, "D": 26, "E": 60}, header_row=header_row)
    sheet["A1"].font = Font(name=FONT, size=14, bold=True)
    sheet["A2"].font = NOTE_FONT


def verdict_summary(verdict: list[dict], findings: list[dict]) -> list[str]:
    lines = []
    for verdict_name in ("Keep visible", "Keep in Advanced", "Candidate to remove"):
        names = [v["setting"] for v in verdict if v["verdict"] == verdict_name]
        if names:
            lines.append(f"{verdict_name}: {', '.join(names)}.")
    high = [f["id"] + " " + f["finding"] for f in findings if f["severity"] == "High"]
    if high:
        lines.append("Fix first: " + " | ".join(high))
    return lines


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
